from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Slide 1: Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Director Cut AI"
    subtitle.text = "Fully Automated, Local AI Video Editing Pipeline\n\nDeveloped by Aniket Tripathi"

    # Slide 2: The Problem
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Problem"
    tf = slide.placeholders[1].text_frame
    tf.text = "Video editing has a steep learning curve and high technical barriers."
    p = tf.add_paragraph()
    p.text = "Manual Editing is Tedious: Finding engaging moments requires frame-by-frame scrubbing."
    p = tf.add_paragraph()
    p.text = "Privacy Risks: Current AI tools force users to upload raw media to cloud APIs."
    p = tf.add_paragraph()
    p.text = "Costly Dependencies: Subscriptions for both editing software and AI generations."

    # Slide 3: The Solution
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Solution: Director Cut AI"
    tf = slide.placeholders[1].text_frame
    tf.text = "An autonomous, locally-hosted AI video director."
    p = tf.add_paragraph()
    p.text = "Zero-Touch Pipeline: Upload footage and let AI analyze motion and audio automatically."
    p = tf.add_paragraph()
    p.text = "100% Local Execution: Powered entirely by on-device LLMs and neural networks."
    p = tf.add_paragraph()
    p.text = "Generative Audio: Synthesizes context-aware human voiceovers dynamically."

    # Slide 4: System Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "System Architecture"
    tf = slide.placeholders[1].text_frame
    tf.text = "A robust Python/React pipeline utilizing distinct AI engines:"
    p = tf.add_paragraph()
    p.text = "1. Faster-Whisper: Extracts word-level audio timestamps."
    p = tf.add_paragraph()
    p.text = "2. OpenCV: Analyzes frame-by-frame visual energy."
    p = tf.add_paragraph()
    p.text = "3. Llama 3.1: Acts as the director to pick the best 5-second slice and write a script."
    p = tf.add_paragraph()
    p.text = "4. Kokoro-82M TTS: Generates the synthetic voiceover."
    p = tf.add_paragraph()
    p.text = "5. MoviePy v2.0: Cuts, ducks background audio, layers TTS, and exports as .MP4."

    # Slide 5: The Tech Stack
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "The Tech Stack"
    tf = slide.placeholders[1].text_frame
    tf.text = "Frontend: React.js, Vite, custom Glassmorphism CSS architecture."
    p = tf.add_paragraph()
    p.text = "Backend: Python, FastAPI, Uvicorn, Server-Sent Events (SSE)."
    p = tf.add_paragraph()
    p.text = "AI/ML Models: Llama 3.1 (Ollama), Whisper, Kokoro TTS."
    p = tf.add_paragraph()
    p.text = "Media Processing: OpenCV, MoviePy."

    # Slide 6: Technical Challenges Overcome
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Engineering Challenges Overcome"
    tf = slide.placeholders[1].text_frame
    tf.text = "Browser Caching: Defeated rigid browser caching using dynamic timestamp URL generation."
    p = tf.add_paragraph()
    p.text = "Autoplay Restrictions: Engineered strict H.264/AAC MP4 encoding to bypass silent browser blocks."
    p = tf.add_paragraph()
    p.text = "Data Safeguards: Built mathematical fallback functions to prevent crashing on videos lacking audio tracks."

    # Slide 7: About the Developer
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "About the Developer"
    tf = slide.placeholders[1].text_frame
    tf.text = "Aniket Tripathi"
    
    p = tf.add_paragraph()
    p.text = "B.Tech in Computer Science and Engineering, Institute of Engineering and Technology (IET) Lucknow (Class of 2027)."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Competitive Programming Achievements:"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "800+ algorithmic problems solved across 100+ contests."
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Codeforces: AIR 1849 | CodeChef: Global Rank 4663."
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Portfolio: portfolio-iota-black-81.vercel.app"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Contact: tripathianiket2005@gmail.com"
    p.level = 1

    # Save the file
    prs.save("Director_Cut_AI_Presentation.pptx")
    print("Success! Your presentation has been generated: Director_Cut_AI_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()